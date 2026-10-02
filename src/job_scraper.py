# -*- coding: utf-8 -*-
"""
job_scraper.py  -  เก็บประกาศงานจาก JobBKK และ Jobthai แล้ว Export เป็น .xlsx

ตัวแปรที่เก็บ (ตามไฟล์ "ตัวแปรที่ต้องการจัดเก็บ"):
job_title | salary/THB | company_location | company_size | years_experience |
experience_level | education_level | company_name | industry | remote_status |
employment_type | ข้อมูลจาก

หลักการ:
  - ตรวจ robots.txt ทุก URL ก่อนดึง (urllib.robotparser) และเว้นระยะระหว่าง request
  - เก็บเฉพาะประกาศงาน ไม่เก็บ resume / อีเมล / เบอร์โทร (ไม่ยุ่งข้อมูลส่วนบุคคล)
  - ใช้ CONFIG ด้านล่างปรับ keyword, จำนวนหน้า, URL ค้นหา
"""
import json
import logging
import random
import re
import time
import urllib.robotparser as robotparser
from datetime import datetime
from urllib.parse import quote_plus, urljoin, urlparse

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ============================ CONFIG ============================
KEYWORDS = ["data"]
MAX_PAGES = 3                 # จำนวนหน้าผลค้นหาต่อ keyword
MAX_JOBS_PER_SITE = 100      # เพดานจำนวนประกาศต่อเว็บ
DELAY_SECONDS = 10            # เว้นระยะระหว่าง request (จะใช้ค่าที่มากกว่า Crawl-delay ใน robots.txt)
FETCH_COMPANY_PAGE = True     # เปิดหน้าบริษัทเพื่อหา company_size / industry ถ้าหน้างานไม่มี
OUTPUT_FILE = "job_data.xlsx"
SAVE_EVERY = 20               # บันทึกไฟล์ทุก N ประกาศ กันข้อมูลหายถ้าโปรแกรมหยุดกลางทาง

# ใส่อีเมลติดต่อจริงของคุณ เพื่อให้เว็บไซต์ติดต่อกลับได้ (มารยาทของ bot)
USER_AGENT = "StudentResearchBot/1.0 (Khon Kaen University data-science project; patipat.ho@kkumail.com)"

# !!! ตรวจสอบ search_url ก่อนใช้จริง: เปิดเว็บ ค้นหางาน แล้วคัดลอก URL ผลลัพธ์มาแทนที่
#     {kw} = คำค้น, {page} = เลขหน้า  (ถ้า URL เว็บไม่มี paging แบบนี้ให้ปรับตามจริง)
SITES = {
    "JobBKK": {
        "search_url": "https://www.jobbkk.com/jobs/lists/{page}/%E0%B8%AB%E0%B8%B2%E0%B8%87%E0%B8%B2%E0%B8%99,%E0%B8%97%E0%B8%B1%E0%B9%89%E0%B8%87%E0%B8%AB%E0%B8%A1%E0%B8%94,%E0%B8%97%E0%B8%B8%E0%B8%81%E0%B8%88%E0%B8%B1%E0%B8%87%E0%B8%AB%E0%B8%A7%E0%B8%B1%E0%B8%94,%E0%B8%97%E0%B8%B1%E0%B9%89%E0%B8%87%E0%B8%AB%E0%B8%A1%E0%B8%94.html?occupationText=%E0%B9%84%E0%B8%AD%E0%B8%97%E0%B8%B5%20/%20(Data%20Engineer/Big%20Data,%20Data%20Scientist,%20Database%20Administration)&lang_id=1&occupation=25-397,25-210,25-222",   # <-- ตัวอย่าง ต้องตรวจสอบ
        "detail_re": re.compile(r"/jobs/detail(?:urgent)?/\d+/\d+"),   # detailurgent = หน้าประกาศแบบเดียวกัน (หน้าผลลัพธ์หมวดนี้ใช้ลิงก์ชนิดนี้ทั้งหมด)
        # สำรอง: ถ้าไม่เจอใน <a> ให้ดึงคู่เลข /jobs/detail/<company>/<job> จาก HTML/JSON ดิบ
        "id_re": re.compile(r"[\\]?/jobs[\\]?/detail(?:urgent)?[\\]?/(\d+)[\\]?/(\d+)"),
        "id_url": "https://www.jobbkk.com/jobs/detail/{0}/{1}",
    },
    "Jobthai": {
        "search_url": "https://www.jobthai.com/%E0%B8%AB%E0%B8%B2%E0%B8%87%E0%B8%B2%E0%B8%99/%E0%B8%87%E0%B8%B2%E0%B8%99%E0%B8%84%E0%B8%AD%E0%B8%A1%E0%B8%9E%E0%B8%B4%E0%B8%A7%E0%B9%80%E0%B8%95%E0%B8%AD%E0%B8%A3%E0%B9%8C-%E0%B9%84%E0%B8%AD%E0%B8%97%E0%B8%B5,%E0%B8%87%E0%B8%B2%E0%B8%99-data-scientist-data-engineer-data-analyst?page={page}",  # <-- ตัวอย่าง ต้องตรวจสอบ
        "detail_re": re.compile(r"/(?:th|en)/(?:company/)?job/\d+(?!/preview)"),
        # ถ้าไม่เจอลิงก์ใน <a> จะดึงเลข id จาก HTML/JSON ดิบแล้วสร้าง URL ตามรูปแบบนี้
        "id_re": re.compile(r"[\\]?/(?:th|en)[\\]?/(?:company[\\]?/)?job[\\]?/(\d+)"),
        "id_url": "https://www.jobthai.com/th/job/{0}",
        # หน้า /company/job/ID เป็น layout แบบแบรนด์บริษัท ข้อมูลไม่ครบ -> ใช้หน้าปกติ /th/job/ID แทน
        "canonical": (r"/(?:th|en)/company/job/", "/th/job/"),
    },
}

COLUMNS = [
    "job_title", "salary/THB", "company_location", "company_size",
    "years_experience", "experience_level", "education_level", "company_name",
    "industry", "remote_status", "employment_type", "ข้อมูลจาก",
]
# ================================================================

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("scraper")

session = requests.Session()
session.headers.update({"User-Agent": USER_AGENT, "Accept-Language": "th,en;q=0.8"})

# ----------------------- label ที่ใช้ค้นหาในหน้าเว็บ -----------------------
LABELS = {
    "location": ["location", "สถานที่ปฏิบัติงาน", "สถานที่ทำงาน", "สถานที่"],
    "salary": ["salary", "เงินเดือน (บาท)", "เงินเดือน"],
    "employment_type": ["รูปแบบงาน", "job type", "employment type", "ประเภทงาน"],
    "level": ["ระดับตำแหน่งงาน", "job level", "position level"],
    "company_size": ["company size", "number of employees", "employees", "ขนาดบริษัท", "จำนวนพนักงาน"],
    "industry": ["business type", "industry", "ประเภทธุรกิจ", "ประเภทกิจการ"],
    "education": ["ระดับการศึกษา", "education level", "education"],
    "experience": ["ประสบการณ์(ปี)", "ประสบการณ์ (ปี)", "ประสบการณ์ทำงาน", "ประสบการณ์",
                   "years of experience", "experience"],
}
ALL_LABELS = {l for v in LABELS.values() for l in v} | {
    "vacancies", "job descriptions", "qualifications", "จำนวนที่รับ", "อัตรา", "คุณสมบัติ",
    "รายละเอียดงาน", "คุณสมบัติผู้สมัคร", "ติดต่อ", "contact"}

TRANSIT = re.compile(r"^(bts|mrt|brt|airport-link)[\s-]*icon", re.I)

# ข้อความไอคอนที่อาจติดมากับ text ของ Jobthai/JobBKK
ICON_PREFIX = re.compile(
    r"^(pin location|salary icon|number of positions icon|hybrid work icon|"
    r"online interview icon|bts icon|mrt icon|brt icon|airport-link-icon)\s*", re.I)


# ============================ robots.txt ============================
class Robots:
    def __init__(self):
        self.cache = {}

    def _get(self, url):
        p = urlparse(url)
        base = f"{p.scheme}://{p.netloc}"
        if base not in self.cache:
            rp = robotparser.RobotFileParser()
            try:
                r = session.get(base + "/robots.txt", timeout=20)
                if r.status_code == 200:
                    rp.parse(r.text.splitlines())
                else:
                    rp.allow_all = True
            except requests.RequestException:
                rp.disallow_all = True      # ถ้าอ่าน robots.txt ไม่ได้ ให้ปลอดภัยไว้ก่อน
            self.cache[base] = rp
            log.info("loaded robots.txt: %s", base)
        return self.cache[base]

    def allowed(self, url):
        return self._get(url).can_fetch(USER_AGENT, url)

    def delay(self, url):
        cd = self._get(url).crawl_delay(USER_AGENT)
        return max(DELAY_SECONDS, cd or 0)


robots = Robots()


def fetch(url, retries=3):
    """ดึง HTML อย่างสุภาพ: เช็ก robots.txt, เว้นระยะ, retry เมื่อ error ชั่วคราว"""
    if not robots.allowed(url):
        log.warning("ถูกห้ามโดย robots.txt ข้าม: %s", url)
        return None
    for attempt in range(1, retries + 1):
        time.sleep(robots.delay(url) + random.uniform(0, 3))
        try:
            r = session.get(url, timeout=30)
            if r.status_code == 200:
                r.encoding = r.apparent_encoding or "utf-8"
                return r.text
            if r.status_code in (429, 503):
                wait = 60 * attempt
                log.warning("HTTP %s รอ %ss: %s", r.status_code, wait, url)
                time.sleep(wait)
                continue
            log.warning("HTTP %s ข้าม: %s", r.status_code, url)
            return None
        except requests.RequestException as e:
            log.warning("request error (%s/%s): %s", attempt, retries, e)
    return None


# ============================ ตัวช่วย parse ============================
def get_lines(soup):
    for t in soup(["script", "style", "noscript"]):
        t.decompose()
    lines = []
    for raw in soup.get_text("\n").splitlines():
        raw = raw.strip()
        if TRANSIT.match(raw):          # สถานีรถไฟฟ้า ไม่ใช่ที่อยู่ที่ทำงาน
            continue
        s = ICON_PREFIX.sub("", raw).strip()
        if s:
            lines.append(s)
    return lines


END_MARKERS = {"ติดต่อ", "contact", "เลือกวิธีการสมัครงาน", "สนใจสมัครงานตำแหน่งงานนี้กรุณาติดต่อ"}
COMPANY_HINT = re.compile(
    r"จำกัด|บริษัท|ห้างหุ้นส่วน|co\.?,?\s*ltd|company limited|\bltd\b|\binc\b|corporation|\bpcl\b", re.I)
GENERIC_LINK_TEXT = re.compile(r"งานอื่น|other jobs|ดูทั้งหมด|view all|profile|โปรไฟล์|company story|ดูบริษัท", re.I)


def body_lines(lines):
    """ตัดส่วนท้าย (ช่องทางติดต่อ/ปุ่มสมัคร/footer) ออก เพื่อไม่ให้ regex ไปเจอคำจากเมนูหรือ footer"""
    for i, l in enumerate(lines):
        if l.strip().lower() in END_MARKERS:
            return lines[:i]
    return lines


def company_from_contact(lines):
    """สำรอง: Jobthai ใส่ชื่อบริษัทไว้ใต้หัวข้อ 'ติดต่อ'"""
    for i, l in enumerate(lines):
        if l.strip().lower() in ("ติดต่อ", "contact"):
            for c in lines[i + 1:i + 6]:
                if COMPANY_HINT.search(c) and "@" not in c:
                    return c
    return None


def find_value(lines, labels, look=4):
    """หาค่าหลัง label ทั้งแบบ 'label : value' (บรรทัดเดียว) และ label อยู่บรรทัดหนึ่ง ค่าอยู่บรรทัดถัดไป"""
    for i, line in enumerate(lines):
        low = line.lower().strip()
        for lab in labels:
            if low.rstrip(":： ").strip() == lab:
                for nxt in lines[i + 1:i + 1 + look]:
                    v = nxt.strip(" :：")
                    if (nxt.rstrip().endswith((":", "：")) or v.lower() in ALL_LABELS
                            or GENERIC_LINK_TEXT.search(v)):
                        break               # เจอ label ถัดไป = label นี้ไม่มีค่า
                    if v:
                        return v
            elif re.match(rf"^{re.escape(lab)}\s*[:：]\s*\S", low):
                return line.split(":", 1)[-1].split("：", 1)[-1].strip()
    return None


def parse_salary(text):
    """คืน (min, max, midpoint) หน่วยบาท; ถ้า 'ตามตกลง/Negotiable' คืน None"""
    if not text:
        return None, None, None
    nums = [int(n.replace(",", "")) for n in re.findall(r"\d{1,3}(?:,\d{3})+|\d{4,}", text)]
    if not nums:
        return None, None, None
    lo, hi = nums[0], nums[1] if len(nums) > 1 else nums[0]
    return lo, hi, (lo + hi) / 2


def parse_years(text, labeled=None):
    """คืน (จำนวนปีขั้นต่ำ, แหล่งที่มา) แหล่งที่มา = 'label' | 'text' | None
    - ใช้ค่าจาก label ก่อน (เช่น JobBKK 'ประสบการณ์(ปี) : 4 ปีขึ้นไป')
    - ถ้าไม่มี/ไม่มีตัวเลข ค่อยหาจากข้อความประกาศ (ใกล้คำว่า ประสบการณ์/experience)
    - 'ไม่ระบุ' = ไม่มีข้อมูล (ว่าง) ไม่ใช่ 0 ; 0 ใช้เฉพาะกรณีประกาศบอกว่าไม่ต้องมีประสบการณ์/จบใหม่
    - รองรับหน่วยเดือน (6 เดือน -> 0.5)"""
    def months(t):
        m = re.search(r"(\d{1,2})\s*(?:เดือน|months?)", t)
        return round(int(m.group(1)) / 12, 2) if m else None

    if labeled:
        lab = labeled.lower()
        if re.search(r"ไม่จำเป็น|ไม่มี|ไม่จำกัด|fresh|จบใหม่", lab):
            return 0, "label"
        m = re.search(r"\d{1,2}", lab)
        if m:
            mo = months(lab)
            return (mo if mo is not None else int(m.group(0))), "label"
    t = text.lower()
    if re.search(r"ไม่มีประสบการณ์|ไม่จำเป็นต้องมีประสบการณ์|fresh\s*grad|จบใหม่|no experience", t):
        return 0, "text"
    pats = [
        r"(?:ประสบการณ์|experience)[^\n]{0,40}?(\d{1,2})(?:\s*[-–~]\s*\d{1,2})?\s*\+?\s*(?:ปี|years?|yrs?)",
        r"(\d{1,2})(?:\s*[-–~]\s*\d{1,2})?\s*\+?\s*(?:ปี|years?|yrs?)[^\n]{0,30}?(?:ประสบการณ์|experience)",
    ]
    for p in pats:
        m = re.search(p, t)
        if m:
            return int(m.group(1)), "text"
    m = re.search(r"(?:ประสบการณ์|experience)[^\n]{0,40}?(\d{1,2})\s*(?:เดือน|months?)", t)
    if m:
        return round(int(m.group(1)) / 12, 2), "text"
    return None, None


def parse_education(text, labeled=None):
    """เลือกวุฒิ 'ต่ำสุด' ที่ปรากฏ (สมมติฐาน: ถือเป็นเกณฑ์ขั้นต่ำ)
    ถ้าหน้าเว็บมี label ระดับการศึกษา (JobBKK) ใช้เฉพาะค่านั้น ไม่ไปกวาดข้อความทั้งหน้า"""
    t = (labeled if labeled else text).lower()
    rules = [  # เรียงจากต่ำไปสูง
        ("High school", r"มัธยม|ม\.\s*[3-6]|high school"),
        ("Diploma/Vocational", r"ปวช|ปวส|อนุปริญญา|diploma|vocational"),
        ("Bachelor", r"ปริญญาตรี|bachelor|\bb\.?sc\b|\bb\.?eng\b"),
        ("Master", r"ปริญญาโท|\bmaster|\bm\.?sc\b"),
        ("PhD", r"ปริญญาเอก|ph\.?\s?d|doctorate"),
    ]
    for name, pat in rules:
        if re.search(pat, t):
            return name
    return None


def parse_level(title, level_text, years):
    blob = f"{title} {level_text or ''}".lower()
    if re.search(r"intern|trainee|ฝึกงาน|junior|entry|fresh|จบใหม่", blob):
        return "Entry/Junior"
    if re.search(r"head|director|manager|หัวหน้า|ผู้จัดการ|principal|lead", blob):
        return "Lead/Manager"
    if re.search(r"senior|sr\.|อาวุโส", blob):
        return "Senior"
    if years is None:
        return None
    if years <= 1:
        return "Entry/Junior"
    if years <= 4:
        return "Mid"
    return "Senior"


def parse_remote(text):
    t = text.lower()
    wfh = re.search(r"work from home|\bwfh\b|\bremote\b|ทำงานที่บ้าน|ทำงานจากที่บ้าน", t)
    if re.search(r"hybrid|ไฮบริด", t) or (wfh and re.search(r"\boffice\b|ออฟฟิศ|เข้าออฟฟิศ", t)):
        return "Hybrid"
    if re.search(r"work from home|\bwfh\b|\bremote\b|ทำงานที่บ้าน|ทำงานจากที่บ้าน", t):
        return "Remote"
    return "Not specified"


def parse_employment(title, labeled):
    """ใช้ค่าจาก label (JobBKK) ถ้าไม่มีให้ดูจากชื่อตำแหน่งเท่านั้น ไม่กวาดทั้งหน้า"""
    src = (labeled or title or "").lower()
    for name, pat in [
        ("Internship", r"intern|ฝึกงาน|สหกิจ"),
        ("Part-time", r"part[\s-]?time|พาร์ท?ไทม์|พาร์ทไทม์"),
        ("Contract", r"contract|สัญญาจ้าง|งานชั่วคราว|freelance|ฟรีแลนซ์"),
        ("Full-time", r"full[\s-]?time|งานประจำ|พนักงานประจำ"),
    ]:
        if re.search(pat, src):
            return name
    return None


def parse_jsonld(soup):
    """ถ้าหน้ามี schema.org JobPosting จะดึงค่าที่เชื่อถือได้มาเติม"""
    out = {}
    for tag in soup.find_all("script", type="application/ld+json"):
        try:
            data = json.loads(tag.string or "")
        except (ValueError, TypeError):
            continue
        for item in (data if isinstance(data, list) else [data]):
            if not isinstance(item, dict) or item.get("@type") != "JobPosting":
                continue
            out["title"] = item.get("title")
            org = item.get("hiringOrganization") or {}
            out["company"] = org.get("name") if isinstance(org, dict) else None
            loc = item.get("jobLocation")
            loc = loc[0] if isinstance(loc, list) and loc else loc
            if isinstance(loc, dict):
                a = loc.get("address") or {}
                out["location"] = ", ".join(
                    x for x in [a.get("addressLocality"), a.get("addressRegion")] if x) or None
            sal = (item.get("baseSalary") or {}).get("value") or {}
            if isinstance(sal, dict):
                out["salary_min"], out["salary_max"] = sal.get("minValue"), sal.get("maxValue")
            out["employment"] = item.get("employmentType")
            if item.get("jobLocationType") == "TELECOMMUTE":
                out["remote"] = "Remote"
    return out


def company_link(soup, base_url):
    for a in soup.select("a[href*='/company/']"):
        href = a["href"]
        txt = a.get_text(strip=True)
        if "/company/job/" in href or not txt or GENERIC_LINK_TEXT.search(txt):
            continue
        return urljoin(base_url, href), txt
    return None, None


# ============================ parse 1 ประกาศ ============================
def parse_job(html, url, site):
    soup = BeautifulSoup(html, "lxml")
    ld = parse_jsonld(soup)
    comp_url, comp_name = company_link(soup, url)
    lines = get_lines(BeautifulSoup(html, "lxml"))
    text_all = "\n".join(body_lines(lines))

    h1 = soup.find("h1")
    h2 = soup.find("h2")
    title = ld.get("title") or (h1.get_text(strip=True) if h1 else None)

    sal_text = find_value(lines, LABELS["salary"])
    lo, hi, mid = parse_salary(sal_text)
    if mid is None and ld.get("salary_min"):
        lo, hi = ld["salary_min"], ld.get("salary_max") or ld["salary_min"]
        mid = (lo + hi) / 2

    years, years_src = parse_years(text_all, find_value(lines, LABELS["experience"]))
    level_text = find_value(lines, LABELS["level"])
    emp_label = find_value(lines, LABELS["employment_type"])

    row = {
        "job_title": title,
        "salary/THB": mid,
        "company_location": ld.get("location") or find_value(lines, LABELS["location"]),
        "company_size": find_value(lines, LABELS["company_size"]),
        "years_experience": years,
        "experience_level": parse_level(title or "", level_text, years),
        "education_level": parse_education(text_all, find_value(lines, LABELS["education"])),
        "company_name": (ld.get("company") or comp_name or company_from_contact(lines)
                         or (h2.get_text(strip=True) if site == "JobBKK" and h2 else None)),
        "industry": find_value(lines, LABELS["industry"]),
        "remote_status": ld.get("remote") or parse_remote(text_all),
        "employment_type": parse_employment(title, emp_label or ld.get("employment")),
        "ข้อมูลจาก": site,
    }
    extra = {"job_url": url, "salary_min": lo, "salary_max": hi,
             "salary_raw": sal_text, "years_source": years_src or "not found", "scraped_at": datetime.now().strftime("%Y-%m-%d %H:%M")}
    return row, extra, comp_url


def enrich_from_company(row, comp_url, cache):
    """เติม company_size / industry จากหน้าบริษัท (ถ้าหน้างานไม่มี)"""
    if not comp_url or (row["company_size"] and row["industry"]):
        return
    if comp_url not in cache:
        html = fetch(comp_url)
        cache[comp_url] = get_lines(BeautifulSoup(html, "lxml")) if html else []
    lines = cache[comp_url]
    row["company_size"] = row["company_size"] or find_value(lines, LABELS["company_size"])
    row["industry"] = row["industry"] or find_value(lines, LABELS["industry"])


# ============================ เก็บลิงก์ประกาศ ============================
def collect_links(site, cfg):
    links, seen = [], set()
    for kw in KEYWORDS:
        for page in range(1, MAX_PAGES + 1):
            url = cfg["search_url"].format(kw=quote_plus(kw), page=page)
            log.info("[%s] search '%s' page %s", site, kw, page)
            html = fetch(url)
            if not html:
                break
            soup = BeautifulSoup(html, "lxml")
            new = 0
            for a in soup.select("a[href]"):
                if cfg["detail_re"].search(a["href"]):
                    full = urljoin(url, a["href"]).split("#")[0].split("?")[0]
                    if cfg.get("canonical"):
                        full = re.sub(cfg["canonical"][0], cfg["canonical"][1], full)
                    if full not in seen:
                        seen.add(full)
                        links.append(full)
                        new += 1
            # สำรอง: หน้าที่วาดด้วย JavaScript มักมี id ประกาศอยู่ใน JSON ดิบ (เช่น __NEXT_DATA__)
            if new == 0 and cfg.get("id_re"):
                for ids in dict.fromkeys(cfg["id_re"].findall(html)):
                    full = cfg["id_url"].format(*(ids if isinstance(ids, tuple) else (ids,)))
                    if full not in seen:
                        seen.add(full)
                        links.append(full)
                        new += 1
            if new == 0:
                dump = f"debug_{site}_page{page}.html"
                with open(dump, "w", encoding="utf-8") as f:
                    f.write(html)
                log.warning("[%s] ไม่พบลิงก์ประกาศ (HTML ยาว %s ตัวอักษร) บันทึกไว้ที่ %s", site, len(html), dump)
                break
            if len(links) >= MAX_JOBS_PER_SITE:
                return links[:MAX_JOBS_PER_SITE]
    return links


# ============================ Export Excel ============================
def export_excel(rows, extras, path=OUTPUT_FILE):
    df = pd.DataFrame(rows, columns=COLUMNS)
    df_extra = pd.DataFrame(extras)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="jobs", index=False)
        df_extra.to_excel(xw, sheet_name="traceability", index=False)
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.font = Font(name="Arial", bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F4E78")
                c.alignment = Alignment(horizontal="center", vertical="center")
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:200])
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(12, width + 2), 50)
                for c in col[1:]:
                    c.font = Font(name="Arial", size=10)
    log.info("บันทึก %s แถว -> %s", len(df), path)


# ============================ main ============================
def main():
    rows, extras, company_cache = [], [], {}
    for site, cfg in SITES.items():
        links = collect_links(site, cfg)
        log.info("[%s] พบลิงก์ประกาศ %s รายการ", site, len(links))
        for i, link in enumerate(links, 1):
            html = fetch(link)
            if not html:
                continue
            try:
                row, extra, comp_url = parse_job(html, link, site)
                if FETCH_COMPANY_PAGE:
                    enrich_from_company(row, comp_url, company_cache)
            except Exception as e:           # ประกาศเดียวพังไม่ควรทำให้ทั้งงานหยุด
                log.exception("parse ผิดพลาด %s: %s", link, e)
                continue
            rows.append(row)
            extras.append(extra)
            if i % SAVE_EVERY == 0:
                export_excel(rows, extras)
    export_excel(rows, extras)


def debug_page(url):
    """python job_scraper.py debug <URL>  -> ดูว่าหน้านั้นมีลิงก์/ข้อมูลอะไรให้ดึงได้"""
    html = fetch(url)
    if not html:
        print("ดึงหน้าไม่ได้ (ถูกห้ามโดย robots.txt หรือ HTTP error)")
        return
    with open("debug_page.html", "w", encoding="utf-8") as f:
        f.write(html)
    soup = BeautifulSoup(html, "lxml")
    hrefs = [a["href"] for a in soup.select("a[href]")]
    print("ความยาว HTML:", len(html))
    print("จำนวนแท็ก <a>:", len(hrefs))
    print("มี __NEXT_DATA__:", "__NEXT_DATA__" in html)
    print("ตัวอย่าง href ที่มีคำว่า job:")
    for h in [h for h in hrefs if "job" in h.lower()][:15]:
        print("  ", h)
    print("บันทึก HTML ไว้ที่ debug_page.html")


def debug_job(url):
    """python job_scraper.py debugjob <URL ประกาศงาน> -> ดูว่าหน้านั้นพูดถึงประสบการณ์ไว้อย่างไร และสคริปต์อ่านได้เท่าไร"""
    html = fetch(url)
    if not html:
        print("ดึงหน้าไม่ได้")
        return
    site = "JobBKK" if "jobbkk" in url else "Jobthai"
    row, extra, _ = parse_job(html, url, site)
    print("ผลที่สคริปต์อ่านได้:")
    for k in ("job_title", "years_experience", "experience_level", "education_level"):
        print(f"  {k}: {row[k]}")
    print("  years_source:", extra["years_source"])
    key = re.compile(r"ประสบการณ์|experience|years?|yrs?|ปี|เดือน", re.I)
    print("\nบรรทัดในหน้าที่เกี่ยวกับประสบการณ์/ปี:")
    for line in get_lines(BeautifulSoup(html, "lxml")):
        if key.search(line) and "@" not in line and len(line) < 250:
            print("  ", line)


if __name__ == "__main__":
    import sys
    if len(sys.argv) >= 3 and sys.argv[1] == "debugjob":
        debug_job(sys.argv[2])
    elif len(sys.argv) >= 3 and sys.argv[1] == "debug":
        debug_page(sys.argv[2])
    else:
        main()
